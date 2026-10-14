"""The Hyper Detail sheet must not materialise empty cells.

WHY THIS TEST EXISTS. On 2026-10-14 the payments pull began taking the GitHub runner down
— not failing the step, killing the VM, so that 28 consecutive runs pulled the data
correctly and could not commit it, and payments-automation/data went 12 hours stale. The
risk monitor, which reads transactions.csv, went degraded with it.

The cause was this sheet. It is 56,077 rows x ~348 columns = ~19.5 MILLION positions, and
openpyxl allocates a full Cell object for a None exactly as it does for a value. The fee
matrix is a sparse join — roughly 503,765 fee rows spread across those positions, so ~97%
of it is empty. Writing the empties cost about 4.6 GB of cells holding nothing.

Measured at full production size (56,077 x 348, 2.6% filled):

    writing every cell   ~19.5M cells   ~4.6 GB
    skipping the Nones    1.88M cells    614 MB

A WARNING FOR WHOEVER COMES NEXT. The obvious suspect in refresh_workbook.py is the
copy.copy() of font/fill/alignment/border per cell. It is NOT the problem: openpyxl
de-duplicates styles into a shared table on assignment, so the copy is discarded
immediately. Measured peak RSS is identical with and without it. It wastes CPU, not memory.
"""
import io

import openpyxl
import pytest
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROWS, BASE, PLAN_COLS = 400, 25, 60
FILL_EVERY = 7          # ~14% filled; the real sheet is ~2.6%


def _build(skip_empties):
    """A faithful miniature of refresh_hyper_detail's fee matrix."""
    wb = openpyxl.Workbook()
    ws = wb.active
    font = Font(name="Arial", size=8)
    for i in range(ROWS):                       # base columns: real values either way
        for c in range(1, BASE + 1):
            ws.cell(2 + i, c, i * c)
    for c in range(BASE + 1, BASE + PLAN_COLS + 1):
        ws.cell(1, c, "plan%d" % c)             # header keeps the column alive
        ws.column_dimensions[get_column_letter(c)].width = 12
        for i in range(ROWS):
            v = round((i * c) / 100.0, 4) if ((i + c) % FILL_EVERY == 0) else None
            if skip_empties and v is None:
                continue
            cell = ws.cell(2 + i, c, v)
            cell.font = font
            cell.number_format = "#,##0.0000"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return openpyxl.load_workbook(buf).active


@pytest.fixture(scope="module")
def sheets():
    return _build(skip_empties=False), _build(skip_empties=True)


def test_the_grid_is_the_same_shape(sheets):
    """Headers on row 1 keep every column, so nothing narrows when the body is sparse."""
    old, new = sheets
    assert (old.max_row, old.max_column) == (new.max_row, new.max_column)


def test_every_value_is_identical(sheets):
    old, new = sheets
    diffs = [(r, c) for r in range(1, old.max_row + 1)
             for c in range(1, old.max_column + 1)
             if old.cell(r, c).value != new.cell(r, c).value]
    assert not diffs, "values differ at %r" % diffs[:5]


def test_number_formats_survive_where_there_is_a_number(sheets):
    """The format only ever mattered on a populated cell."""
    old, new = sheets
    bad = [(r, c) for r in range(2, old.max_row + 1)
           for c in range(BASE + 1, old.max_column + 1)
           if old.cell(r, c).value is not None
           and old.cell(r, c).number_format != new.cell(r, c).number_format]
    assert not bad, "number format lost at %r" % bad[:5]


def test_column_widths_and_grouping_survive(sheets):
    old, new = sheets
    widths = lambda ws: {k: v.width for k, v in ws.column_dimensions.items() if v.width}
    assert widths(old) == widths(new)


def test_the_fee_matrix_really_is_sparse(sheets):
    """The premise of the change. If the matrix were dense this would buy nothing.

    Counts the FEE MATRIX region only — the base columns to its left are genuinely
    populated and are written either way.
    """
    _old, new_ws = sheets
    positions = ROWS * PLAN_COLS
    populated = sum(1 for r in range(2, new_ws.max_row + 1)
                    for c in range(BASE + 1, new_ws.max_column + 1)
                    if new_ws.cell(r, c).value is not None)
    ratio = populated / positions
    # This miniature is ~1/7 filled by construction; production measured ~2.6%.
    assert ratio < 0.25, (
        "fee matrix is %.1f%% populated — if it is no longer sparse, skipping empty "
        "cells buys nothing and the memory problem needs a different fix" % (100 * ratio))
